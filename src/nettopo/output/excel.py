"""Documentação em planilha Excel (inventário, interfaces, links, VLANs, IPs, rotas...)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from nettopo.models import Topology
from nettopo.output.icons import ROLE_LABELS, TIER_LABELS

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _status(v: bool | None) -> str:
    return "up" if v is True else "down" if v is False else ""


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        width = max([len(str(headers[c - 1]))] + [len(str(r[c - 1] or "")) for r in rows[:500]]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 8), 60)
    ws.freeze_panes = "A2"
    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        tname = "T_" + "".join(ch for ch in title if ch.isalnum())[:20]
        table = Table(displayName=tname, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)


def write_excel(topo: Topology, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    m = topo.meta
    s = m.stats
    ws["A1"] = m.name or "Topologia de rede"
    ws["A1"].font = Font(bold=True, size=16)
    info = [
        ("Sementes", ", ".join(m.seeds)), ("Início", m.started_at.strftime("%d/%m/%Y %H:%M:%S")),
        ("Fim", m.finished_at.strftime("%d/%m/%Y %H:%M:%S") if m.finished_at else ""),
        ("Duração (s)", s.get("duration_s", "")), ("Equipamentos coletados", s.get("collected", 0)),
        ("Vistos apenas por vizinhos", s.get("stubs", 0)), ("Links físicos", s.get("physical_links", 0)),
        ("Links L3", s.get("l3_links", 0)), ("Adjacências lógicas", s.get("logical_links", 0)),
        ("Interfaces", s.get("interfaces", 0)), ("Alvos com falha", s.get("failed_targets", 0)),
        ("Versão NetTopo", m.tool_version),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    row = len(info) + 5
    ws.cell(row=row, column=1, value="Por fabricante").font = Font(bold=True, size=12)
    for j, (k, v) in enumerate(s.get("by_vendor", {}).items(), start=row + 1):
        ws.cell(row=j, column=1, value=k)
        ws.cell(row=j, column=2, value=v)
    ws.cell(row=row, column=4, value="Por papel").font = Font(bold=True, size=12)
    for j, (k, v) in enumerate(s.get("by_role", {}).items(), start=row + 1):
        ws.cell(row=j, column=4, value=ROLE_LABELS.get(k, k))
        ws.cell(row=j, column=5, value=v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["D"].width = 22

    devs = sorted(topo.devices.values(), key=lambda d: (d.stub, (d.label or "").lower()))
    name = {d.id: d.label for d in devs}

    _sheet(wb, "Inventário",
           ["Hostname", "IP gerência", "Papel", "Camada", "Fabricante", "Sistema", "Versão", "Modelo", "Serial",
            "Uptime", "Local", "Coleta", "Interfaces", "Vizinhos", "Descoberto por", "Profundidade", "Descrição"],
           [[d.label, d.mgmt_ip, ROLE_LABELS.get(d.role, d.role), TIER_LABELS.get(d.tier or "", d.tier), d.vendor,
             d.os, d.os_version, d.model, d.serial, d.uptime, d.location,
             "visto por vizinho" if d.stub else ", ".join(d.collected_via), len(d.interfaces), len(d.neighbors),
             name.get(d.discovered_from or "", ""), d.depth if not d.stub else "", (d.sys_description or "")[:250]]
            for d in devs])

    _sheet(wb, "Links",
           ["ID", "Origem", "Interface origem", "Destino", "Interface destino", "Tipo", "Protocolos", "Velocidade (Mbps)",
            "Sub-rede", "LAG"],
           [[lk.id, name.get(lk.source), lk.source_interface, name.get(lk.target), lk.target_interface, lk.kind,
             ", ".join(lk.protocols), lk.speed_mbps, lk.subnet, lk.lag] for lk in topo.links])

    _sheet(wb, "Interfaces",
           ["Equipamento", "Interface", "Descrição", "IPv4", "Admin", "Oper", "Velocidade (Mbps)", "MTU", "MAC",
            "LAG", "VLAN acesso", "Modo", "VRF"],
           [[d.label, i.name, i.description, " ".join(i.ipv4), _status(i.admin_up), _status(i.oper_up), i.speed_mbps,
             i.mtu, i.mac, i.parent, i.access_vlan, i.mode, i.vrf]
            for d in devs if not d.stub for i in d.interfaces])

    _sheet(wb, "Vizinhos",
           ["Equipamento", "Protocolo", "Interface local", "Vizinho", "Interface remota", "IP vizinho", "Plataforma",
            "Capacidades", "ASN", "Estado"],
           [[d.label, n.protocol, n.local_interface, n.remote_hostname, n.remote_interface, n.remote_mgmt_ip,
             n.remote_platform, ", ".join(n.remote_capabilities), n.remote_asn, n.state]
            for d in devs for n in d.neighbors])

    vlan_rows: dict[int, dict[str, Any]] = {}
    for d in devs:
        for v in d.vlans:
            e = vlan_rows.setdefault(v.id, {"names": set(), "devs": []})
            if v.name:
                e["names"].add(v.name)
            e["devs"].append(d.label)
    _sheet(wb, "VLANs", ["VLAN", "Nome(s)", "Qtde equipamentos", "Equipamentos"],
           [[vid, " / ".join(sorted(e["names"])), len(e["devs"]), ", ".join(e["devs"])]
            for vid, e in sorted(vlan_rows.items())])

    _sheet(wb, "Endereçamento IP", ["Endereço", "Equipamento", "Interface", "VRF", "Descrição"],
           [[ip, d.label, i.name, i.vrf, i.description] for d in devs for i in d.interfaces for ip in i.ipv4])

    _sheet(wb, "ARP", ["Equipamento", "IP", "MAC", "Interface"],
           [[d.label, a.ip, a.mac, a.interface] for d in devs for a in d.arp])

    _sheet(wb, "Rotas", ["Equipamento", "Prefixo", "Next-hop", "Interface", "Protocolo", "VRF"],
           [[d.label, r.prefix, r.next_hop, r.interface, r.protocol, r.vrf] for d in devs for r in d.routes])

    if any(d.mac_table for d in devs):
        _sheet(wb, "Tabela MAC", ["Equipamento", "MAC", "VLAN", "Interface"],
               [[d.label, e.mac, e.vlan, e.interface] for d in devs for e in d.mac_table])

    _sheet(wb, "Falhas", ["Alvo", "Motivo"], [[k, v] for k, v in sorted(m.failed_targets.items())])
    wb.save(path)
    return path
