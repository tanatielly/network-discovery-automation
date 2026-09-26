"""Pipeline completo (descoberta -> grafo -> saídas) usando a rede simulada."""

import asyncio
import json
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nettopo.collectors.simulated import SimulatedCollector, demo_settings
from nettopo.discovery.engine import discover
from nettopo.models import Topology
from nettopo.output import export_all
from nettopo.topology.diff import diff_topologies


@pytest.fixture(scope="module")
def topo() -> Topology:
    sim = SimulatedCollector()
    return asyncio.run(discover(demo_settings(sim), collector=sim))


def _link(topo, a, b):
    return [lk for lk in topo.links if {lk.source, lk.target} == {a, b}]


def test_discovers_whole_lab(topo):
    collected = {d.hostname for d in topo.devices.values() if not d.stub}
    assert len(collected) == 16
    assert {"EDGE-RTR-01", "BRANCH-RTR-01", "BRANCH-SW-01", "SRV-01", "ACC-SW-05"} <= collected
    stubs = {d.hostname for d in topo.devices.values() if d.stub}
    assert stubs == {"ISP-A-PE", "ISP-B-PE", "AP-01", "AP-02", "SEP001122334455"}
    assert set(topo.meta.failed_targets) == {"10.0.3.21", "10.0.3.22"}


def test_no_duplicate_devices_or_links(topo):
    hosts = [d.hostname for d in topo.devices.values()]
    assert len(hosts) == len(set(hosts))
    # 24 links do laboratório + 1 link L3 (filial via OSPF)
    assert len([lk for lk in topo.links if lk.kind == "physical"]) == 24
    assert len(_link(topo, "core-sw-01", "core-sw-02")) == 2  # dois links paralelos (LAG)


def test_bidirectional_links_merged_with_both_interfaces(topo):
    lk = _link(topo, "edge-rtr-01", "fw-01")[0]
    ends = {(lk.source, lk.source_interface), (lk.target, lk.target_interface)}
    assert ends == {("edge-rtr-01", "GigabitEthernet0/0/1"), ("fw-01", "port1")}
    assert lk.subnet == "10.255.0.0/30"
    assert set(_link(topo, "core-sw-02", "dist-sw-02")[0].protocols) == {"lldp", "cdp"}


def test_l3_and_routing_enrichment(topo):
    branch = _link(topo, "edge-rtr-01", "branch-rtr-01")
    assert len(branch) == 1 and branch[0].kind == "l3" and "ospf" in branch[0].protocols
    assert "bgp" in _link(topo, "edge-rtr-01", "stub-isp-a-pe")[0].protocols
    assert all(lk.lag for lk in _link(topo, "core-sw-01", "core-sw-02"))


def test_roles_and_tiers(topo):
    by = {d.hostname: d for d in topo.devices.values()}
    assert by["EDGE-RTR-01"].role == "router" and by["EDGE-RTR-01"].tier == "edge"
    assert by["FW-01"].role == "firewall" and by["FW-02"].tier == "security"
    assert by["CORE-SW-01"].tier == "core" and by["CORE-SW-02"].tier == "core"
    assert by["DIST-SW-01"].tier == "distribution" and by["DIST-SW-02"].tier == "distribution"
    assert {by[h].tier for h in ("ACC-SW-01", "ACC-SW-02", "ACC-SW-03", "ACC-SW-04", "ACC-SW-05")} == {"access"}
    assert by["AP-01"].role == "ap" and by["SEP001122334455"].role == "phone"
    assert by["ISP-A-PE"].tier == "external"
    assert by["SRV-01"].role == "server"


def test_exports(topo, tmp_path: Path):
    paths = export_all(topo, ["html", "drawio", "xlsx", "md", "json"], tmp_path, "lab")
    html, drawio, xlsx, md, js = paths
    text = html.read_text(encoding="utf-8")
    assert "cytoscape" in text and "EDGE-RTR-01" in text and "</script>" in text
    root = ET.parse(drawio).getroot()
    assert root.tag == "mxfile" and len(root.findall("diagram")) == 2
    wb = load_workbook(xlsx)
    assert {"Resumo", "Inventário", "Links", "Interfaces", "VLANs", "Rotas"} <= set(wb.sheetnames)
    assert wb["Inventário"].max_row == len(topo.devices) + 1
    readme = md.read_text(encoding="utf-8")
    assert "```mermaid" in readme and (md.parent / "devices" / "edge-rtr-01.md").exists()
    assert Topology.load(str(js)).meta.stats == json.loads(js.read_text(encoding="utf-8"))["meta"]["stats"]


def test_diff(topo):
    new = topo.model_copy(deep=True)
    new.devices["core-sw-01"].os_version = "4.32.0F"
    removed = new.links.pop(0)
    d = diff_topologies(topo, new)
    assert d["devices_changed"][0]["changes"]["os_version"] == ("4.31.2F", "4.32.0F")
    assert len(d["links_removed"]) == 1 and removed.source in d["links_removed"][0]


def test_each_device_collected_once_under_concurrency():
    """Mesmo equipamento anunciado por LLDP e OSPF (IPs diferentes) não é coletado duas vezes."""
    sim = SimulatedCollector(delay=0.05)
    done: list[str] = []

    def progress(ev):
        if ev["type"] == "device_done":
            done.append(ev["hostname"])

    asyncio.run(discover(demo_settings(sim, concurrency=4), collector=sim, progress=progress))
    assert len(done) == len(set(done)) == 16
